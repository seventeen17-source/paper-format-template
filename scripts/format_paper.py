#!/usr/bin/env python3
"""Apply the distilled paper-format-template rules to an existing DOCX.

This script changes formatting only. It does not intentionally rewrite text.
It uses conservative heuristics for semantic roles and preserves ambiguous content.
"""

from __future__ import annotations

import argparse
import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

A4_W = Cm(21.0)
A4_H = Cm(29.7)
MANUSCRIPT_MARGINS = dict(top=Cm(2.299), bottom=Cm(1.700), left=Cm(2.501), right=Cm(1.501))
INDENT = Cm(0.741)

HEADING_RE = re.compile(r"^\s*\d+(?:\.\d+){0,2}\s+\S")
FIG_RE = re.compile(r"^\s*(?:Fig\.|Figure|图)\s*\d+[\.:。]?", re.I)
TABLE_RE = re.compile(r"^\s*(?:Table|表)\s*\d+[\.:。]?", re.I)
ABSTRACT_RE = re.compile(r"^\s*(?:Abstract\s*:|摘要\s*[：:])", re.I)
KEYWORDS_RE = re.compile(r"^\s*(?:Keywords?\s*:|关键词\s*[：:])", re.I)
REFERENCES_RE = re.compile(r"^\s*(?:References|参考文献)\s*$", re.I)
ARTICLE_TYPE_RE = re.compile(r"^\s*(?:Original papers?|Research article|Article|原创论文)\s*$", re.I)


def set_rfonts(run, latin: str, east_asia: str) -> None:
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east_asia)


def format_runs(paragraph, size: float, bold=None, latin="Times New Roman", east_asia="SimSun") -> None:
    for run in paragraph.runs:
        set_rfonts(run, latin, east_asia)
        run.font.name = latin
        run.font.size = Pt(size)
        if bold is not None:
            run.bold = bold


def set_para(paragraph, *, align, line, left=INDENT, right=INDENT, first=INDENT, before=0, after=0) -> None:
    paragraph.alignment = align
    pf = paragraph.paragraph_format
    pf.line_spacing = line
    pf.left_indent = left
    pf.right_indent = right
    pf.first_line_indent = first
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)


def has_drawing(paragraph) -> bool:
    return bool(paragraph._p.xpath(".//w:drawing|.//w:pict"))


def bold_prefix(paragraph, pattern: re.Pattern) -> None:
    """Bold only a matched text prefix without changing visible text."""
    text = paragraph.text
    m = pattern.match(text)
    if not m:
        return
    prefix_len = m.end()
    seen = 0
    for run in list(paragraph.runs):
        run_text = run.text
        run_len = len(run_text)
        start, end = seen, seen + run_len
        if end <= prefix_len:
            run.bold = True
        elif start >= prefix_len:
            run.bold = False
        else:
            cut = prefix_len - start
            left_text, right_text = run_text[:cut], run_text[cut:]
            run.text = left_text
            run.bold = True
            if right_text:
                clone = deepcopy(run._r)
                for child in list(clone):
                    if child.tag != qn("w:rPr"):
                        clone.remove(child)
                t = OxmlElement("w:t")
                t.text = right_text
                clone.append(t)
                clone_rpr = clone.find(qn("w:rPr"))
                if clone_rpr is None:
                    clone_rpr = OxmlElement("w:rPr")
                    clone.insert(0, clone_rpr)
                b = clone_rpr.find(qn("w:b"))
                if b is not None:
                    clone_rpr.remove(b)
                bcs = clone_rpr.find(qn("w:bCs"))
                if bcs is not None:
                    clone_rpr.remove(bcs)
                run._r.addnext(clone)
        seen = end


def set_cell_borders(cell, **edges) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge, attrs in edges.items():
        tag = "w:" + edge
        el = tc_borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            tc_borders.append(el)
        for key, val in attrs.items():
            el.set(qn("w:" + key), str(val))


def set_table_borders(table, **edges) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge, attrs in edges.items():
        tag = "w:" + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        for key, val in attrs.items():
            el.set(qn("w:" + key), str(val))


def is_equation_table(table) -> bool:
    if len(table.rows) != 1 or len(table.columns) != 3:
        return False
    last = table.cell(0, 2).text.strip()
    return bool(re.fullmatch(r"\(\d+\)", last))


def is_abbreviation_table(table) -> bool:
    text = " ".join(c.text for r in table.rows[:2] for c in r.cells)
    return "Abbreviations" in text or "缩写" in text


def format_equation_table(table) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    none = {"val": "none", "sz": "0", "space": "0", "color": "auto"}
    set_table_borders(table, top=none, left=none, bottom=none, right=none, insideH=none, insideV=none)
    for row in table.rows:
        for cell in row.cells:
            set_cell_borders(cell, top=none, left=none, bottom=none, right=none)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    table.cell(0, 1).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    table.cell(0, 2).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for p in table.cell(0, 2).paragraphs:
        format_runs(p, 10.5, latin="Times New Roman", east_asia="SimSun")


def format_data_table(table) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    none = {"val": "none", "sz": "0", "space": "0", "color": "auto"}
    half = {"val": "single", "sz": "4", "space": "0", "color": "auto"}
    one = {"val": "single", "sz": "8", "space": "0", "color": "auto"}
    set_table_borders(table, top=half, left=none, bottom=half, right=none, insideH=none, insideV=none)
    for ri, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_borders(cell, left=none, right=none)
            if ri == 0:
                set_cell_borders(cell, top=one, bottom=one)
            else:
                set_cell_borders(cell, top=none, bottom=none)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.left_indent = Cm(0)
                p.paragraph_format.right_indent = Cm(0)
                format_runs(p, 10.5, latin="Times New Roman", east_asia="SimSun")


def add_line_numbering(section) -> None:
    sect_pr = section._sectPr
    old = sect_pr.find(qn("w:lnNumType"))
    if old is not None:
        sect_pr.remove(old)
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1")
    ln.set(qn("w:restart"), "continuous")
    sect_pr.append(ln)


def format_manuscript(doc: Document) -> None:
    for section in doc.sections:
        section.page_width = A4_W
        section.page_height = A4_H
        section.top_margin = MANUSCRIPT_MARGINS["top"]
        section.bottom_margin = MANUSCRIPT_MARGINS["bottom"]
        section.left_margin = MANUSCRIPT_MARGINS["left"]
        section.right_margin = MANUSCRIPT_MARGINS["right"]
        add_line_numbering(section)

    nonempty = [p for p in doc.paragraphs if p.text.strip()]
    title_element = None
    if nonempty:
        first = nonempty[0]
        if ARTICLE_TYPE_RE.match(first.text.strip()) and len(nonempty) > 1:
            title_element = nonempty[1]._p
        else:
            title_element = first._p

    abstract_index = None
    for i, p in enumerate(doc.paragraphs):
        if ABSTRACT_RE.match(p.text.strip()):
            abstract_index = i
            break

    refs_mode = False
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if not text:
            continue

        if has_drawing(p):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            continue

        if title_element is not None and p._p is title_element:
            set_para(p, align=WD_ALIGN_PARAGRAPH.LEFT, line=1.5, left=INDENT, right=INDENT, first=Cm(0))
            format_runs(p, 18, bold=True, latin="Times New Roman", east_asia="SimHei")
            continue

        if ARTICLE_TYPE_RE.match(text):
            set_para(p, align=WD_ALIGN_PARAGRAPH.LEFT, line=1.5, left=INDENT, right=INDENT, first=Cm(0))
            format_runs(p, 10.5, latin="Times New Roman", east_asia="SimSun")
            continue

        if ABSTRACT_RE.match(text):
            set_para(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line=1.7, left=INDENT, right=INDENT, first=Cm(0))
            format_runs(p, 10.5, latin="Times New Roman", east_asia="SimSun")
            bold_prefix(p, ABSTRACT_RE)
            continue

        if KEYWORDS_RE.match(text):
            set_para(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line=1.7, left=INDENT, right=INDENT, first=Cm(0))
            format_runs(p, 10.5, latin="Times New Roman", east_asia="SimSun")
            bold_prefix(p, KEYWORDS_RE)
            continue

        if REFERENCES_RE.match(text):
            refs_mode = True
            set_para(p, align=WD_ALIGN_PARAGRAPH.LEFT, line=2.0, left=INDENT, right=INDENT, first=Cm(0), before=15.6, after=15.6)
            format_runs(p, 12, bold=True, latin="Times New Roman", east_asia="SimHei")
            continue

        if HEADING_RE.match(text):
            refs_mode = False
            set_para(p, align=WD_ALIGN_PARAGRAPH.LEFT, line=2.0, left=INDENT, right=INDENT, first=Cm(0), before=15.6, after=15.6)
            format_runs(p, 12, bold=True, latin="Times New Roman", east_asia="SimHei")
            continue

        if FIG_RE.match(text):
            align = WD_ALIGN_PARAGRAPH.CENTER if len(text) <= 120 else WD_ALIGN_PARAGRAPH.JUSTIFY
            set_para(p, align=align, line=2.0, left=Cm(0), right=Cm(0), first=Cm(0))
            format_runs(p, 10.5, latin="Times New Roman", east_asia="SimSun")
            bold_prefix(p, FIG_RE)
            continue

        if TABLE_RE.match(text):
            set_para(p, align=WD_ALIGN_PARAGRAPH.CENTER, line=2.0, left=Cm(0), right=Cm(0), first=Cm(0))
            format_runs(p, 10.5, latin="Times New Roman", east_asia="SimSun")
            bold_prefix(p, TABLE_RE)
            continue

        if refs_mode:
            set_para(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line=2.0, left=INDENT, right=INDENT, first=Cm(0))
            format_runs(p, 10.5, latin="Times New Roman", east_asia="SimSun")
            continue

        if abstract_index is not None and i < abstract_index:
            set_para(p, align=WD_ALIGN_PARAGRAPH.LEFT, line=1.5, left=INDENT, right=INDENT, first=Cm(0))
            format_runs(p, 11, latin="Times New Roman", east_asia="SimSun")
            continue

        set_para(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line=2.0, left=INDENT, right=INDENT, first=INDENT)
        format_runs(p, 10.5, latin="Times New Roman", east_asia="SimSun")

    for table in doc.tables:
        if is_equation_table(table):
            format_equation_table(table)
        elif is_abbreviation_table(table):
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        format_runs(p, 10.5, latin="Times New Roman", east_asia="SimSun")
        else:
            format_data_table(table)


def main() -> None:
    ap = argparse.ArgumentParser(description="Apply paper-format-template formatting without rewriting content.")
    ap.add_argument("input", type=Path)
    ap.add_argument("output", type=Path)
    args = ap.parse_args()
    doc = Document(args.input)
    format_manuscript(doc)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    print(args.output)


if __name__ == "__main__":
    main()
