#!/usr/bin/env python3
"""Lightweight format audit for manuscript DOCX files."""
from __future__ import annotations
import argparse
from pathlib import Path
from docx import Document


def cm(v):
    return None if v is None else round(v.cm, 3)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("docx", type=Path)
    args = ap.parse_args()
    d = Document(args.docx)
    s = d.sections[0]
    print("page_cm", cm(s.page_width), cm(s.page_height))
    print("margins_cm", cm(s.top_margin), cm(s.bottom_margin), cm(s.left_margin), cm(s.right_margin))
    print("paragraphs", len(d.paragraphs), "tables", len(d.tables))
    for i, p in enumerate(d.paragraphs):
        t = p.text.strip()
        if t:
            print(i, p.style.name, repr(t[:80]), "align", p.alignment, "line", p.paragraph_format.line_spacing)


if __name__ == "__main__":
    main()
